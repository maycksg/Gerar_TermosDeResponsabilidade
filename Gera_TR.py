import os
import re
import pandas as pd
from docx import Document

# 1. Definir o diretório base dinamicamente (pasta onde o script está sendo executado)
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Caminhos dinâmicos para a planilha, o modelo e a pasta de saída
EXCEL_PATH = os.path.join(BASE_DIR, 'Base.xlsx')
MODELO_PATH = os.path.join(BASE_DIR, 'MODELO.docx')
OUTPUT_DIR = os.path.join(BASE_DIR, 'Termos_Gerados')

def limpar_nome_arquivo(nome: str) -> str:
    """
    Remove ou substitui caracteres inválidos para nomes de arquivos no Windows/Linux/macOS.
    """
    # Remove caracteres proibidos: \ / : * ? " < > |
    nome_limpo = re.sub(r'[\\/*?:"<>|]', '', nome)
    # Remove espaços extras no início e no fim
    return nome_limpo.strip()

# 2. Verificar se a planilha e o modelo existem antes de continuar
if not os.path.exists(EXCEL_PATH):
    raise FileNotFoundError(f"Arquivo Excel não encontrado: {EXCEL_PATH}")

if not os.path.exists(MODELO_PATH):
    raise FileNotFoundError(f"Modelo Word não encontrado: {MODELO_PATH}")

# 3. Carregar a planilha do Excel, tratando CPF e CNPJ como strings
df = pd.read_excel(EXCEL_PATH, dtype={'Cpf': str, 'Cnpj': str})

# Substituir valores nulos por strings vazias e remover espaços
df['Cpf'] = df['Cpf'].fillna('').str.strip()
df['Cnpj'] = df['Cnpj'].fillna('').str.strip()
df['Filial'] = df['Filial'].fillna('Desconhecida').str.strip()

# 4. Iterar sobre cada linha da planilha
for index, row in df.iterrows():
    nome = str(row.get('Nome', '')).strip()
    cargo = str(row.get('Cargo', '')).strip()
    cpf = row.get('Cpf', '')
    cnpj = row.get('Cnpj', '')
    modelo = str(row.get('Modelo', '')).strip()
    tag = str(row.get('Tag', '')).strip()
    tipo = str(row.get('Tipo', '')).strip()
    filial = str(row.get('Filial', 'Desconhecida')).strip()
    estado = str(row.get('Estado', '')).strip()
    endereco = str(row.get('Endereco', '')).strip()
    numero = str(row.get('Numero', '')).strip()
    bairro = str(row.get('Bairro', '')).strip()
    cep = str(row.get('CEP', '')).strip()

    # Se o nome estiver em branco, pula a linha ou define um padrão
    if not nome:
        print(f"Linha {index + 2}: Nome em branco. Pulando...")
        continue

    # Limpar o nome para salvar no sistema de arquivos
    nome_sanitizado = limpar_nome_arquivo(nome)
    filial_sanitizada = limpar_nome_arquivo(filial)

    # Pasta de destino da filial dentro do diretório de saída
    pasta_filial = os.path.join(OUTPUT_DIR, filial_sanitizada)

    try:
        os.makedirs(pasta_filial, exist_ok=True)
    except PermissionError:
        print(f"Aviso: Sem permissão para criar a pasta: {pasta_filial}")
        continue

    # Carregar o documento modelo do Word
    doc = Document(MODELO_PATH)

    # Dicionário de substituições
    substituicoes = {
        '<<Nome>>': nome,
        '<<Cargo>>': cargo,
        '<<Cpf>>': cpf,
        '<<Cnpj>>': cnpj,
        '<<Modelo>>': modelo,
        '<<Tag>>': tag,
        '<<Tipo>>': tipo,
        '<<Filial>>': filial,
        '<<Estado>>': estado,
        '<<Endereco>>': endereco,
        '<<Numero>>': numero,
        '<<Bairro>>': bairro,
        '<<CEP>>': cep
    }

    # Substituir no corpo do texto (parágrafos)
    for paragraph in doc.paragraphs:
        for key, value in substituicoes.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)

    # Substituir também dentro de tabelas (se houver)
    for table in doc.tables:
        for row_table in table.rows:
            for cell in row_table.cells:
                for paragraph in cell.paragraphs:
                    for key, value in substituicoes.items():
                        if key in paragraph.text:
                            paragraph.text = paragraph.text.replace(key, value)

    # 5. Nome do arquivo dinâmico e padronizado
    nome_arquivo = f"Termo_de_Responsabilidade_{nome_sanitizado}.docx"
    caminho_arquivo = os.path.join(pasta_filial, nome_arquivo)

    # Salvar o documento final
    doc.save(caminho_arquivo)
    print(f"Gerado: {caminho_arquivo}")

print("\nTodos os documentos foram gerados com sucesso! 🎉")